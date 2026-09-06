"""Approval-bound DOCX rendering and optional local PDF conversion."""

import asyncio
import hashlib
import os
import subprocess
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Protocol
from uuid import uuid4

from docx import Document

from app.ports.local_backend import ArtifactStore, DraftResolver, StoredArtifact, StoredDraft
from app.storage.session_workspace import LocalSessionWorkspaceStore, WorkspaceArea
from app.tools.contracts import (
    ArtifactFormat,
    ArtifactReference,
    DocumentExportExecutionRequest,
    DocumentExportResult,
)
from app.workflow.contracts import ExecutionStatus


class PdfConversionError(RuntimeError):
    """Local conversion failed without invalidating the rendered DOCX."""


class PdfConverter(Protocol):
    async def convert(self, source: Path, output_directory: Path) -> Path: ...


class LibreOfficePdfConverter:
    """Invoke one configured local LibreOffice executable without a shell."""

    def __init__(self, executable: str = "soffice", *, timeout_seconds: float = 60) -> None:
        if not executable.strip() or timeout_seconds <= 0:
            raise ValueError("converter executable and timeout must be valid")
        self._executable = executable
        self._timeout = timeout_seconds

    async def convert(self, source: Path, output_directory: Path) -> Path:
        expected = output_directory / f"{source.stem}.pdf"
        if expected.is_symlink() or expected.exists():
            raise PdfConversionError("converter output path already exists")
        try:
            completed = await asyncio.to_thread(
                subprocess.run,
                [
                    self._executable,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(output_directory),
                    str(source),
                ],
                shell=False,
                capture_output=True,
                timeout=self._timeout,
                check=False,
            )
        except (FileNotFoundError, subprocess.TimeoutExpired, OSError) as error:
            raise PdfConversionError("local PDF converter unavailable") from error
        if completed.returncode != 0:
            raise PdfConversionError("local PDF conversion failed")
        if expected.is_symlink() or not expected.is_file():
            raise PdfConversionError("local PDF converter produced no regular output")
        data = expected.read_bytes()
        if not data.startswith(b"%PDF-"):
            expected.unlink(missing_ok=True)
            raise PdfConversionError("local PDF converter produced invalid output")
        return expected


class LocalDocumentArtifactExecutor:
    """Resolve an approved draft, render files, and register exact metadata."""

    def __init__(
        self,
        resolver: DraftResolver,
        artifacts: ArtifactStore,
        workspaces: LocalSessionWorkspaceStore,
        converter: PdfConverter | None = None,
    ) -> None:
        self._resolver = resolver
        self._artifacts = artifacts
        self._workspaces = workspaces
        self._converter = converter

    async def create_artifacts(
        self, request: DocumentExportExecutionRequest
    ) -> DocumentExportResult:
        stored = await self._resolver.resolve_for_export(request)
        if stored is None:
            return self._failure("draft_not_authorized")
        workspace = self._workspaces.get_session_workspace(str(stored.session_id))
        base = f"approval-note-{stored.draft_id}"
        docx_path = self._workspaces.file_path(
            str(stored.session_id), WorkspaceArea.ARTIFACTS, f"{base}.docx"
        )
        created_paths: list[Path] = []
        try:
            await asyncio.to_thread(self._render_docx_atomic, stored, docx_path)
            created_paths.append(docx_path)
            paths: dict[ArtifactFormat, Path] = {ArtifactFormat.DOCX: docx_path}
            if ArtifactFormat.PDF in request.arguments.formats:
                if self._converter is None:
                    raise PdfConversionError("local PDF converter is not configured")
                pdf_path = await self._converter.convert(docx_path, workspace.artifacts)
                self._validate_pdf(pdf_path, workspace.artifacts)
                paths[ArtifactFormat.PDF] = pdf_path
                created_paths.append(pdf_path)

            references: list[ArtifactReference] = []
            for format_ in request.arguments.formats:
                path = paths[format_]
                metadata = self._metadata(request, stored, format_, path)
                await self._artifacts.create(
                    metadata, execution_claim_token=request.execution_claim_token
                )
                references.append(
                    ArtifactReference(
                        artifact_id=metadata.artifact_id,
                        format=format_,
                        file_name=metadata.file_name,
                    )
                )
            if ArtifactFormat.DOCX not in request.arguments.formats:
                docx_path.unlink(missing_ok=True)
            return DocumentExportResult(
                status=ExecutionStatus.COMPLETED, artifacts=tuple(references)
            )
        except PdfConversionError:
            for path in created_paths:
                if path.suffix.lower() == ".pdf":
                    path.unlink(missing_ok=True)
            return self._failure("pdf_conversion_failed")
        except Exception:
            for path in created_paths:
                path.unlink(missing_ok=True)
            return self._failure("artifact_creation_failed")

    @staticmethod
    def _render_docx_atomic(stored: StoredDraft, destination: Path) -> None:
        if destination.exists() or destination.is_symlink():
            raise FileExistsError("artifact destination already exists")
        descriptor, name = tempfile.mkstemp(dir=destination.parent, suffix=".tmp")
        os.close(descriptor)
        temporary = Path(name)
        try:
            document = Document()
            document.add_heading(stored.draft.subject, 0)
            document.add_heading("Executive Summary", level=1)
            document.add_paragraph(stored.draft.summary)
            document.add_heading("Findings", level=1)
            for finding in stored.draft.findings:
                document.add_heading(finding.title, level=2)
                document.add_paragraph(finding.description)
                document.add_paragraph(f"Severity: {finding.severity.value}")
                if finding.uncertainty:
                    document.add_paragraph(f"Uncertainty: {finding.uncertainty}")
            document.add_heading("Recommendations", level=1)
            document.add_paragraph(stored.draft.recommendation)
            document.add_heading("Critical Claims / Evidence", level=1)
            for claim in stored.draft.critical_claims:
                document.add_paragraph(claim.text, style="List Bullet")
                document.add_paragraph("Sources: " + ", ".join(claim.evidence_source_ids))
            document.add_heading("Uncertainties / Limitations", level=1)
            for uncertainty in stored.draft.uncertainties:
                document.add_paragraph(uncertainty, style="List Bullet")
            document.add_heading("Source References", level=1)
            for source_id in stored.draft.evidence_source_ids:
                document.add_paragraph(source_id, style="List Bullet")
            document.save(str(temporary))
            with temporary.open("ab") as stream:
                os.fsync(stream.fileno())
            os.link(temporary, destination, follow_symlinks=False)
        finally:
            temporary.unlink(missing_ok=True)

    @staticmethod
    def _validate_pdf(path: Path, root: Path) -> None:
        if path.is_symlink() or path.resolve(strict=True).parent != root.resolve():
            raise PdfConversionError("PDF output escaped the artifact directory")
        if not path.read_bytes().startswith(b"%PDF-"):
            raise PdfConversionError("PDF output is invalid")

    @staticmethod
    def _metadata(
        request: DocumentExportExecutionRequest,
        stored: StoredDraft,
        format_: ArtifactFormat,
        path: Path,
    ) -> StoredArtifact:
        data = path.read_bytes()
        return StoredArtifact(
            artifact_id=uuid4(),
            session_id=stored.session_id,
            workflow_run_id=stored.workflow_run_id,
            owner_user_id=stored.owner_user_id,
            approval_id=request.approval_id,
            draft_id=stored.draft_id,
            format=format_,
            file_name=path.name,
            size_bytes=len(data),
            sha256=hashlib.sha256(data).hexdigest(),
            created_at=datetime.now(UTC),
        )

    @staticmethod
    def _failure(code: str) -> DocumentExportResult:
        return DocumentExportResult(status=ExecutionStatus.FAILED, failure_code=code)
